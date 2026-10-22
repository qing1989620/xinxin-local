from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import re
from pathlib import Path

from .skillnet.index import BM25Index
from .store import clean_text, now, uid

MAX_UPLOAD = 256 * 1024 * 1024
MAX_EXTRACT = 16 * 1024 * 1024


class FileTooLarge(ValueError):
    pass


class UploadConflict(ValueError):
    pass


TEXT_EXT = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".py",
    ".ts",
    ".tsx",
    ".js",
    ".yaml",
    ".yml",
    ".log",
}


class FileLibrary:
    def __init__(self, store):
        self.store = store
        self.root = store.directory / "files"
        self.root.mkdir(exist_ok=True)
        self.previews = store.directory / "previews"
        self._repair_paths()

    def _repair_paths(self):
        """数据目录被移动/复制后，把指向旧位置的绝对路径改回当前库。

        本地版是从云端项目整目录复制来的，`files.path` 里可能残留上一个
        项目的绝对路径。文件实体按 `id + 后缀` 命名一并复制过来了，因此
        只要当前库里有同名文件，就可以安全地把记录修回本库路径。
        逐条幂等：已经在库内且存在的记录不会被动。
        """
        root = self.root.resolve()
        fixes = []
        for row in self.store.all("SELECT id,path FROM files"):
            stored = Path(row["path"])
            try:
                inside = stored.resolve().is_relative_to(root)
            except OSError:
                inside = False
            if inside and stored.is_file():
                continue
            local = root / stored.name
            if local.is_file():
                fixes.append((str(local), row["id"]))
        if fixes:
            with self.store.connect(write=True) as db:
                db.executemany("UPDATE files SET path=? WHERE id=?", fixes)

    def save(self, name, content, mime=None, run_id=None, extract=True):
        return self.save_stream(name, io.BytesIO(content), mime, run_id, extract)

    def save_stream(
        self, name, source, mime=None, run_id=None, extract=True, request_id=None
    ):
        """Bounded disk copy; a retry key commits with the resource in one transaction."""
        name = clean_text(Path(name.replace("\\", "/")).name[:180])
        if not name or name in {".", ".."}:
            raise ValueError("文件为空")
        mime = mime or mimetypes.guess_type(name)[0] or "application/octet-stream"
        ident = uid()
        suffix = Path(name).suffix.lower()
        suffix = suffix if re.fullmatch(r"\.[a-z0-9]{1,12}", suffix) else ""
        target = self.root / (ident + suffix)
        staging = self.root / (ident + ".incoming")
        total, digest = 0, hashlib.sha256()
        committed = False
        try:
            with staging.open("xb") as output:
                while chunk := source.read(1024 * 1024):
                    total += len(chunk)
                    if total > MAX_UPLOAD:
                        raise FileTooLarge("文件超过 256 MB 上限")
                    output.write(chunk)
                    digest.update(chunk)
            if not total:
                raise ValueError("文件为空")
            checksum = digest.hexdigest()
            signature = json.dumps([name, mime, checksum], ensure_ascii=False)
            text, parse_error, extraction = "", "", {}
            if extract and suffix in {".pdf", ".docx"}:
                from .document_text import extract_document
                try:
                    text, extraction = extract_document(staging, suffix=suffix)
                    parse_error = extraction.get("warning", "")
                except Exception:  # optional parsers cannot discard the uploaded original
                    parse_error = "原文件已保存，正文提取失败；可在学习中重新识别"
                    extraction = {"complete": False, "warning": parse_error}
            elif extract and total <= MAX_EXTRACT:
                try:
                    from .document_text import MAX_TEXT_CHARS
                    text = self.extract(name, staging.read_bytes())
                    if len(text) > MAX_TEXT_CHARS:
                        text = ""
                        parse_error = "原文件已保存；正文超过识别上限，请拆分资料后学习"
                    extraction = {"complete": not bool(parse_error), "warning": parse_error}
                except Exception:  # noqa: BLE001 -- optional document parsers may raise format-specific errors; retain the original.
                    parse_error = "已保存原文件，暂时无法提取文本"
            elif extract:
                parse_error = "原文件已保存；大文本可在学习中按需识别，其他格式可交给 Agent 处理"
                extraction = {"complete": False, "warning": parse_error}
            if parse_error:
                extraction["warning"] = parse_error
                extraction.setdefault("complete", False)
            text = clean_text(text)
            with self.store.connect(write=True) as db:
                previous = (
                    db.execute(
                        "SELECT file_id,signature FROM uploads WHERE request_id=?",
                        (request_id,),
                    ).fetchone()
                    if request_id
                    else None
                )
                if previous:
                    if previous["signature"] != signature:
                        raise UploadConflict("上传重试标识与文件不匹配，请重新选择文件")
                    ident = previous["file_id"]
                else:
                    staging.replace(target)
                    db.execute(
                        "INSERT INTO files VALUES(?,?,?,?,?,?,?,?,?)",
                        (
                            ident,
                            name,
                            mime,
                            total,
                            checksum,
                            str(target),
                            text,
                            run_id,
                            now(),
                        ),
                    )
                    if request_id:
                        db.execute(
                            "INSERT INTO uploads VALUES(?,?,?)",
                            (request_id, ident, signature),
                        )
                    db.execute("INSERT INTO file_extractions VALUES(?,?)", (ident, json.dumps(extraction, ensure_ascii=False)))
            committed = True
            return self.get(ident)
        finally:
            staging.unlink(missing_ok=True)
            if not committed:
                target.unlink(missing_ok=True)

    @staticmethod
    def extract(name, content):
        return clean_text(FileLibrary._extract(name, content))

    @staticmethod
    def _extract(name, content):
        ext = Path(name).suffix.lower()
        if ext in TEXT_EXT:
            return content.decode("utf-8", errors="replace")
        if ext == ".pdf":
            from pypdf import PdfReader

            return "\n\n".join(
                page.extract_text() or ""
                for page in PdfReader(io.BytesIO(content)).pages
            )
        if ext == ".docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            return "\n".join(
                [p.text for p in document.paragraphs]
                + [
                    " | ".join(cell.text for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
            )
        return ""

    def get(self, ident):
        row = self.store.one("SELECT * FROM files WHERE id=?", (ident,))
        if not row:
            raise ValueError("文件不存在")
        extraction = self.store.one("SELECT data FROM file_extractions WHERE file_id=?", (ident,))
        metadata = json.loads(extraction["data"]) if extraction else {}
        metadata.pop("page_spans", None)
        return {
            key: value for key, value in row.items() if key not in {"path", "text"}
        } | {"indexed": bool(row["text"]), "extraction": metadata, "parse_error": metadata.get("warning", "")}

    def path(self, ident):
        row = self.store.one("SELECT * FROM files WHERE id=?", (ident,))
        if not row:
            raise ValueError("文件不存在")
        root = self.root.resolve()
        stored = Path(row["path"])
        # 先信数据库里的路径；若它指向别处（历史数据、目录搬家），
        # 再按文件名在当前库内找一次——始终不越过库根目录。
        for candidate in (stored, root / stored.name):
            try:
                resolved = candidate.resolve()
            except OSError:
                continue
            if resolved.is_relative_to(root) and resolved.is_file():
                return resolved
        raise ValueError("文件不可用")

    def preview(self, ident, size=1600):
        """有界预览图：长边不超过 size，JPEG，磁盘缓存按文件 mtime 失效。

        返回缓存文件路径；调用方负责以 image/jpeg 输出。非图片返回 None。
        """
        row = self.store.one("SELECT name,mime,path FROM files WHERE id=?", (ident,))
        if not row or row["mime"] not in {"image/jpeg", "image/png", "image/webp"}:
            return None
        source = self.path(ident)
        from PIL import Image, ImageOps

        # 已经是小尺寸 JPEG 时直接给原图：省一次重编码，也没有质量损失。
        if row["mime"] == "image/jpeg" and source.stat().st_size <= 512 * 1024:
            with Image.open(source) as probe:
                if max(probe.size) <= size:
                    return source
        stamp = int(source.stat().st_mtime)
        self.previews.mkdir(exist_ok=True)
        cache = self.previews / f"{ident}-{size}-{stamp}.jpg"
        if cache.is_file() and cache.stat().st_size > 0:
            return cache
        resample = getattr(Image, "Resampling", Image).LANCZOS
        with Image.open(source) as image:
            image = ImageOps.exif_transpose(image)
            image.thumbnail((size, size), resample)
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            temporary = cache.with_name(cache.name + ".tmp")
            try:
                image.save(temporary, "JPEG", quality=82, optimize=True, progressive=True)
                temporary.replace(cache)
            finally:
                temporary.unlink(missing_ok=True)
        return cache

    def drop_previews(self, ident):
        for stale in self.previews.glob(f"{ident}-*"):
            stale.unlink(missing_ok=True)

    def search(self, query, k=5):
        rows = self.store.all(
            "SELECT id,name,text FROM files WHERE text!='' ORDER BY created DESC"
        )
        chunks = []
        for row in rows:
            text = row["text"]
            for offset in range(0, len(text), 900):
                chunks.append(
                    {
                        "file_id": row["id"],
                        "filename": row["name"],
                        "offset": offset,
                        "text": text[offset : offset + 1100],
                    }
                )
        if not chunks:
            return []
        index = BM25Index()
        index.fit(
            [str(i) for i in range(len(chunks))], [chunk["text"] for chunk in chunks]
        )
        return [
            {**chunks[int(i)], "score": score}
            for i, score in index.search(query, top_k=k)
        ]

    def read(self, ident, offset=0, limit=12000):
        if offset < 0 or not 1 <= limit <= 16000:
            raise ValueError("读取范围无效")
        row = self.store.one("SELECT name,text FROM files WHERE id=?", (ident,))
        if not row:
            raise ValueError("文件不存在")
        end = min(offset + limit, len(row["text"]))
        return {
            "id": ident,
            "name": row["name"],
            "text": row["text"][offset:end],
            "offset": offset,
            "next_offset": end if end < len(row["text"]) else None,
            "total_chars": len(row["text"]),
        }

    def delete(self, ident):
        target = self.path(ident)
        with self.store.connect(write=True) as db:
            db.execute("DELETE FROM uploads WHERE file_id=?", (ident,))
            db.execute("DELETE FROM file_extractions WHERE file_id=?", (ident,))
            db.execute("DELETE FROM file_ocr_pages WHERE file_id=?", (ident,))
            db.execute("DELETE FROM files WHERE id=?", (ident,))
        target.unlink(missing_ok=True)
        self.drop_previews(ident)
